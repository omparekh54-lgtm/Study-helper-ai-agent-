"""Turn uploaded files into clean, location-aware text chunks.

Every chunk remembers where it came from (page / slide / section) so that every
generated question can cite its source.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

SUPPORTED_EXTENSIONS = {"pdf", "docx", "pptx", "txt", "md"}


class ParseError(ValueError):
    """A problem with the uploaded file that the user can fix (shown verbatim in the UI)."""


@dataclass
class Page:
    loc: int
    text: str


@dataclass
class ChunkData:
    idx: int
    loc_start: int
    loc_end: int
    text: str


@dataclass
class ParsedDocument:
    file_type: str
    source_label: str  # page | slide | section
    pages: list[Page]
    chunks: list[ChunkData]

    @property
    def char_count(self) -> int:
        return sum(len(c.text) for c in self.chunks)


# --------------------------------------------------------------------------------------
# Text clean-up helpers
# --------------------------------------------------------------------------------------

_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_MULTI_SPACE = re.compile(r"[ \t ]+")
_MULTI_NEWLINE = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _HYPHEN_BREAK.sub(r"\1\2", text)  # "exam-\nple" -> "example"
    text = _MULTI_SPACE.sub(" ", text)
    lines = [ln.strip() for ln in text.split("\n")]
    # Re-flow hard-wrapped lines into paragraphs, keep blank-line paragraph breaks and bullets.
    paragraphs: list[str] = []
    current: list[str] = []
    for ln in lines:
        if not ln:
            if current:
                paragraphs.append(" ".join(current))
                current = []
            continue
        starts_block = bool(re.match(r"^([\-•●▪*•]|\d+[.)]|[a-zA-Z][.)])\s", ln))
        if current and (starts_block or current[-1].endswith((":", "."))) and len(current[-1]) < 60:
            paragraphs.append(" ".join(current))
            current = []
        current.append(ln)
    if current:
        paragraphs.append(" ".join(current))
    return _MULTI_NEWLINE.sub("\n\n", "\n\n".join(p for p in paragraphs if p.strip())).strip()


def detect_file_type(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("doc", "ppt"):
        raise ParseError(
            f"Old .{ext} files aren't supported. Please save it as .{ext}x (or PDF) and upload again."
        )
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError("Unsupported file type. Upload a PDF, DOCX, PPTX, TXT or Markdown file.")
    if ext == "pdf" and not data.startswith(b"%PDF"):
        raise ParseError("This file doesn't look like a valid PDF.")
    if ext in ("docx", "pptx") and not data.startswith(b"PK"):
        raise ParseError(f"This file doesn't look like a valid .{ext} document.")
    return "txt" if ext == "md" else ext


# --------------------------------------------------------------------------------------
# Format-specific extraction
# --------------------------------------------------------------------------------------


def _extract_pdf(data: bytes, max_pages: int) -> list[Page]:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ParseError("We couldn't open this PDF — it may be corrupted.") from exc
    if doc.needs_pass:
        raise ParseError("This PDF is password-protected. Remove the password and upload again.")
    if doc.page_count > max_pages:
        raise ParseError(f"This PDF has {doc.page_count} pages — the limit is {max_pages}. Try splitting it.")
    pages: list[Page] = []
    for i, page in enumerate(doc):
        # Text blocks ≈ paragraphs/headings, which keeps structure far better than raw lines.
        paragraphs = []
        for block in page.get_text("blocks", sort=True):
            if block[6] != 0:  # 0 = text block, 1 = image
                continue
            text = _HYPHEN_BREAK.sub(r"\1\2", block[4].strip())
            text = _MULTI_SPACE.sub(" ", text.replace("\n", " ")).strip()
            if text:
                paragraphs.append(text)
        pages.append(Page(loc=i + 1, text="\n\n".join(paragraphs)))
    doc.close()
    return pages


def _extract_docx(data: bytes) -> list[Page]:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError("We couldn't open this Word document — it may be corrupted.") from exc

    blocks: list[tuple[bool, str]] = []  # (is_heading, text)
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        is_heading = style.startswith("heading") or style == "title"
        is_list = "list" in style
        blocks.append((is_heading, f"• {text}" if is_list else text))
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append((False, " | ".join(dict.fromkeys(cells))))
    return _group_sections(blocks)


def _extract_pptx(data: bytes, max_pages: int) -> list[Page]:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError("We couldn't open this PowerPoint file — it may be corrupted.") from exc
    if len(prs.slides) > max_pages:
        raise ParseError(f"This deck has {len(prs.slides)} slides — the limit is {max_pages}.")

    pages: list[Page] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        pages.append(Page(loc=i, text=clean_text("\n\n".join(parts))))
    return pages


def _extract_text(data: bytes) -> list[Page]:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 always decodes
        raise ParseError("We couldn't read this text file's encoding.")
    blocks: list[tuple[bool, str]] = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        is_heading = para.startswith("#") or (len(para) < 80 and "\n" not in para and not para.endswith("."))
        blocks.append((is_heading, clean_text(para.lstrip("# "))))
    return _group_sections(blocks)


def _group_sections(blocks: list[tuple[bool, str]], target: int = 1800) -> list[Page]:
    """Group paragraph blocks into numbered 'sections' for formats without pages."""
    sections: list[Page] = []
    current: list[str] = []
    size = 0
    for is_heading, text in blocks:
        if current and (is_heading and size > 400 or size + len(text) > target):
            sections.append(Page(loc=len(sections) + 1, text="\n\n".join(current)))
            current, size = [], 0
        current.append(text)
        size += len(text)
    if current:
        sections.append(Page(loc=len(sections) + 1, text="\n\n".join(current)))
    return sections


# --------------------------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------------------------

_SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(])")


def _split_long(paragraph: str, max_len: int) -> list[str]:
    if len(paragraph) <= max_len:
        return [paragraph]
    out, buf = [], ""
    for sentence in _SENTENCE_SPLIT.split(paragraph):
        if buf and len(buf) + len(sentence) + 1 > max_len:
            out.append(buf)
            buf = sentence
        else:
            buf = f"{buf} {sentence}".strip()
        while len(buf) > max_len:  # a single monster "sentence"
            out.append(buf[:max_len])
            buf = buf[max_len:]
    if buf:
        out.append(buf)
    return out


def _is_heading(text: str) -> bool:
    t = text.strip()
    return (
        0 < len(t) < 90
        and len(t.split()) <= 12
        and not t.endswith((".", ",", ";", "?", "!"))
        and not t.startswith(("•", "-", "*", "●"))
    )


def make_chunks(pages: list[Page], target: int = 1400, max_len: int = 2000, min_page: int = 250) -> list[ChunkData]:
    """Group paragraphs into ~target-sized chunks that keep headings with their content."""
    chunks: list[ChunkData] = []
    buf: list[tuple[int, str]] = []  # (loc, paragraph)

    def size() -> int:
        return sum(len(p) + 2 for _, p in buf)

    def flush() -> None:
        nonlocal buf
        carry: list[tuple[int, str]] = []
        while len(buf) > 1 and _is_heading(buf[-1][1]):  # never end a chunk on a heading
            carry.insert(0, buf.pop())
        if buf:
            text = "\n\n".join(p for _, p in buf)
            chunks.append(ChunkData(len(chunks) + 1, buf[0][0], buf[-1][0], text))
        buf = carry

    for page in pages:
        if not page.text.strip():
            continue
        paragraphs = [p for para in page.text.split("\n\n") for p in _split_long(para.strip(), max_len) if p]
        for para in paragraphs:
            current = size()
            if buf and (current + len(para) > target or (_is_heading(para) and current >= target * 0.5)):
                flush()
            buf.append((page.loc, para))
        # Keep chunks aligned to page boundaries unless the page was tiny (e.g. a title slide).
        if size() >= min_page:
            flush()
    while buf:
        before = len(buf)
        flush()
        if len(buf) == before:  # only headings left
            chunks.append(ChunkData(len(chunks) + 1, buf[0][0], buf[-1][0], "\n\n".join(p for _, p in buf)))
            break
    return chunks


# --------------------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------------------


def parse_document(filename: str, data: bytes, *, max_pages: int = 400, max_chars: int = 450_000) -> ParsedDocument:
    if not data:
        raise ParseError("The file is empty.")
    file_type = detect_file_type(filename, data)

    if file_type == "pdf":
        pages, label = _extract_pdf(data, max_pages), "page"
    elif file_type == "docx":
        pages, label = _extract_docx(data), "section"
    elif file_type == "pptx":
        pages, label = _extract_pptx(data, max_pages), "slide"
    else:
        pages, label = _extract_text(data), "section"

    total_chars = sum(len(p.text) for p in pages)
    non_empty = [p for p in pages if len(p.text) > 20]
    if total_chars < 200 or (file_type == "pdf" and len(pages) >= 3 and total_chars / len(pages) < 40):
        raise ParseError(
            "We couldn't find enough readable text. If this is a scanned PDF (photos of pages), "
            "please upload a text-based PDF, DOCX or PPTX instead."
        )
    if total_chars > max_chars:
        raise ParseError(
            f"This document is very long (~{total_chars // 1000}k characters; the limit is "
            f"{max_chars // 1000}k). Split it into chapters and upload them separately."
        )
    chunks = make_chunks(non_empty)
    return ParsedDocument(file_type=file_type, source_label=label, pages=pages, chunks=chunks)


def loc_label(source_label: str, start: int, end: int) -> str:
    short = {"page": ("p.", "pp."), "slide": ("slide", "slides"), "section": ("section", "sections")}
    single, plural = short.get(source_label, ("p.", "pp."))
    return f"{single} {start}" if start == end else f"{plural} {start}–{end}"
