import io

import pytest

from app.services.parsing import ParseError, clean_text, loc_label, make_chunks, parse_document, Page


def test_sample_pdf_parses_into_page_aware_chunks(sample_pdf):
    parsed = parse_document("notes.pdf", sample_pdf)
    assert parsed.file_type == "pdf"
    assert parsed.source_label == "page"
    assert len(parsed.pages) == 4
    assert 5 <= len(parsed.chunks) <= 10
    assert all(c.loc_start <= c.loc_end for c in parsed.chunks)
    assert [c.idx for c in parsed.chunks] == list(range(1, len(parsed.chunks) + 1))
    text = " ".join(c.text for c in parsed.chunks)
    assert "RuBisCO" in text and "photolysis" in text
    # Headings open chunks rather than dangling at the end of the previous one.
    assert any(c.text.startswith("4. Limiting Factors") for c in parsed.chunks)
    assert not any(c.text.rstrip().endswith("4. Limiting Factors") for c in parsed.chunks)


def test_docx_parsing_groups_sections():
    import docx

    d = docx.Document()
    d.add_heading("Cell Biology", level=1)
    for i in range(6):
        d.add_paragraph(f"Mitochondria are the site of aerobic respiration in paragraph {i}. " * 5)
    d.add_heading("Membranes", level=1)
    d.add_paragraph("The plasma membrane is a phospholipid bilayer with embedded proteins. " * 4)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Organelle", "Function"
    table.cell(1, 0).text, table.cell(1, 1).text = "Ribosome", "Protein synthesis"
    buf = io.BytesIO()
    d.save(buf)

    parsed = parse_document("cells.docx", buf.getvalue())
    assert parsed.file_type == "docx" and parsed.source_label == "section"
    joined = " ".join(c.text for c in parsed.chunks)
    assert "phospholipid bilayer" in joined and "Ribosome | Protein synthesis" in joined


def test_pptx_parsing_includes_speaker_notes():
    from pptx import Presentation

    prs = Presentation()
    for title, body, notes in [
        ("Newton's laws", "An object stays at rest unless acted on by a resultant force. " * 3, "Mention inertia examples."),
        ("Momentum", "Momentum equals mass multiplied by velocity and is conserved in collisions. " * 3, "Solve a worked example."),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)

    parsed = parse_document("physics.pptx", buf.getvalue())
    assert parsed.source_label == "slide"
    joined = " ".join(c.text for c in parsed.chunks)
    assert "Speaker notes: Solve a worked example." in joined


def test_markdown_is_treated_as_text():
    md = "# Algebra\n\n" + ("Quadratic equations have at most two real roots. " * 10) + "\n\n## Factorising\n\n" + ("Look for common factors first. " * 10)
    parsed = parse_document("algebra.md", md.encode())
    assert parsed.file_type == "txt"
    assert "Quadratic equations" in parsed.chunks[0].text


@pytest.mark.parametrize(
    "filename,data,message",
    [
        ("empty.pdf", b"", "empty"),
        ("notes.xyz", b"hello world" * 50, "Unsupported file type"),
        ("old.doc", b"whatever" * 50, "Old .doc"),
        ("fake.pdf", b"not a pdf at all" * 50, "valid PDF"),
        ("fake.docx", b"not a zip" * 50, "valid .docx"),
        ("tiny.txt", b"Too short.", "readable text"),
    ],
)
def test_bad_files_get_friendly_errors(filename, data, message):
    with pytest.raises(ParseError, match=message):
        parse_document(filename, data)


def test_scanned_pdf_is_rejected():
    import pymupdf

    doc = pymupdf.open()
    for _ in range(4):
        doc.new_page()
    with pytest.raises(ParseError, match="scanned"):
        parse_document("scan.pdf", doc.tobytes())


def test_char_limit():
    text = ("Photosynthesis converts light energy into chemical energy. " * 200).encode()
    with pytest.raises(ParseError, match="very long"):
        parse_document("long.txt", text, max_chars=5_000)


def test_clean_text_fixes_hyphenation_and_reflows():
    assert clean_text("photo-\nsynthesis is\ngreat.") == "photosynthesis is great."


def test_make_chunks_never_ends_with_heading():
    pages = [Page(1, "Intro\n\n" + "A" * 900 + ".\n\nNext Heading\n\n" + "B" * 900 + ".")]
    chunks = make_chunks(pages, target=1000)
    assert chunks[1].text.startswith("Next Heading")


def test_loc_labels():
    assert loc_label("page", 3, 3) == "p. 3"
    assert loc_label("page", 3, 5) == "pp. 3–5"
    assert loc_label("slide", 2, 2) == "slide 2"
